"""
Script para generar documentos de prueba realistas para el módulo de Ingesta Documental y NLP.
Crea dos escenarios corporativos complementarios:
1. Documento PDF: Auditoría de Ciberseguridad, Continuidad y Riesgos (Sector Salud / IPS).
2. Documento DOCX: Diagnóstico de Gobernanza, Gestión de Cambios y Proveedores (Sector Logística).
"""

import os
from pathlib import Path
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
import docx

DATA_DIR = Path(__file__).parent

def generar_pdf_salud():
    pdf_path = DATA_DIR / "Auditoria_Ciberseguridad_y_Continuidad_Hospital_San_Rafael.pdf"
    doc = SimpleDocTemplate(
        str(pdf_path),
        pagesize=letter,
        rightMargin=45,
        leftMargin=45,
        topMargin=45,
        bottomMargin=45
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#0F172A'),
        alignment=1,
        spaceAfter=6
    )
    subtitle_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#475569'),
        alignment=1,
        spaceAfter=15
    )
    h2_style = ParagraphStyle(
        'Heading2_Custom',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#1E3A8A'),
        spaceBefore=12,
        spaceAfter=6
    )
    body_style = ParagraphStyle(
        'Body_Custom',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#334155'),
        spaceAfter=6
    )
    table_cell = ParagraphStyle(
        'TableCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=11,
        textColor=colors.HexColor('#1E293B')
    )
    table_header = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=11,
        textColor=colors.white
    )

    elements = []

    # Encabezado
    elements.append(Paragraph("HOSPITAL UNIVERSITARIO SAN RAFAEL E.S.E.", title_style))
    elements.append(Paragraph("INFORME DE AUDITORÍA TÉCNICA: CIBERSEGURIDAD, CONTINUIDAD Y GESTIÓN DE RIESGOS DE TI", subtitle_style))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#1E3A8A'), spaceAfter=12))

    # Ficha Técnica
    ficha_data = [
        [Paragraph("<b>Entidad Evaluada:</b> Hospital San Rafael", table_cell), Paragraph("<b>Sector:</b> Salud y Asistencia Médica", table_cell)],
        [Paragraph("<b>Tipo de Auditoría:</b> Técnica y Regulatoria", table_cell), Paragraph("<b>Fecha de Emisión:</b> Octubre 2026", table_cell)],
        [Paragraph("<b>Normativas de Referencia:</b> COBIT 2019 / Circular 007 Supersalud", table_cell), Paragraph("<b>Clasificación:</b> Confidencial", table_cell)]
    ]
    t_ficha = Table(ficha_data, colWidths=[260, 260])
    t_ficha.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F8FAFC')),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
    ]))
    elements.append(t_ficha)
    elements.append(Spacer(1, 10))

    # 1. Introducción y Antecedentes
    elements.append(Paragraph("1. Contexto y Alcance de la Revisión", h2_style))
    elements.append(Paragraph(
        "El presente informe recopila los hallazgos identificados en la infraestructura tecnológica del Hospital San Rafael. "
        "La entidad atiende a más de 250,000 usuarios y custodia historias clínicas digitales con información sensible de salud. "
        "La auditoría evaluó los controles de seguridad de la información, procedimientos de copias de seguridad (backup), "
        "gestión de incidentes de servicio en urgencias y cumplimiento de la regulación de protección de datos personales.",
        body_style
    ))

    # 2. Situaciones Críticas y Problemas Detectados
    elements.append(Paragraph("2. Problemas y Vulnerabilidades Críticas Identificadas", h2_style))
    elements.append(Paragraph(
        "Durante las pruebas de vulnerabilidad y entrevistas se evidenciaron las siguientes debilidades operacionales:",
        body_style
    ))
    elements.append(Paragraph(
        "• <b>Problema en Ciberseguridad e Identidades:</b> La entidad carece de autenticación de múltiples factores (MFA) en el acceso remoto por VPN y correos electrónicos de los médicos, lo cual genera un alto riesgo de infección por ransomware y exfiltración de historias clínicas.",
        body_style
    ))
    elements.append(Paragraph(
        "• <b>Fallas en Copias de Seguridad (Backup):</b> Se identificó que las copias de seguridad de la base de datos de pacientes no cuentan con pruebas periódicas de restauración, y los respaldos se guardan en el mismo centro de cómputo principal sin sitio de contingencia alterno.",
        body_style
    ))
    elements.append(Paragraph(
        "• <b>Ausencia de Plan de Continuidad (DRP):</b> No existe un Plan de Recuperación ante Desastres (DRP) documentado ni simulacros de continuidad ante caídas masivas del software asistencial en salas de urgencias y cuidados intensivos.",
        body_style
    ))
    elements.append(Paragraph(
        "• <b>Demoras en la Mesa de Ayuda:</b> La mesa de ayuda presenta una tasa de retraso del 45% en la resolución de incidentes de red, afectando la agilidad en la facturación y dispensación de medicamentos.",
        body_style
    ))

    # 3. Necesidades Prioritarias
    elements.append(Paragraph("3. Necesidades Estratégicas y Regulatorias Obligatorias", h2_style))
    elements.append(Paragraph(
        "Para subsanar las no conformidades y mitigar sanciones de la Superintendencia Nacional de Salud, la organización:",
        body_style
    ))
    elements.append(Paragraph(
        "1. <b>Requiere</b> formalizar e implementar de forma prioritaria la Política General de Seguridad de la Información (SGSI) alineada con los procesos COBIT APO13 y DSS05.",
        body_style
    ))
    elements.append(Paragraph(
        "2. <b>Necesita</b> establecer acuerdos de nivel de servicio (SLA) con el proveedor del software de historia clínica para garantizar tiempos máximos de restablecimiento ante incidentes graves.",
        body_style
    ))
    elements.append(Paragraph(
        "3. <b>Debe</b> ejecutar simulacros semestrales de restauración de backups para asegurar la integridad de la base de datos hospitalaria.",
        body_style
    ))
    elements.append(Paragraph(
        "4. <b>Es fundamental</b> capacitar al personal asistencial y administrativo en prevención de ataques de ingeniería social y phishing.",
        body_style
    ))

    # 4. Oportunidades de Mejora y Modernización
    elements.append(Paragraph("4. Oportunidades de Optimización Tecnológica", h2_style))
    elements.append(Paragraph(
        "• <b>Oportunidad de Modernización Cloud:</b> Migrar los respaldos secundarios a una infraestructura en la nube con almacenamiento inmutable para prevenir el secuestro de datos por malware.",
        body_style
    ))
    elements.append(Paragraph(
        "• <b>Automatización del Monitoreo:</b> Implementar herramientas de monitoreo continuo (SIEM) para detectar intrusiones y accesos anómalos en tiempo real.",
        body_style
    ))
    elements.append(Paragraph(
        "• <b>Potencial de Mejora en el Service Desk:</b> Habilitar un portal de autoservicio con base de conocimientos para solucionar bloqueos de contraseñas de médicos en turnos nocturnos.",
        body_style
    ))

    # Tabla Resumen de Procesos COBIT Impactados
    elements.append(Spacer(1, 6))
    elements.append(Paragraph("5. Matriz de Procesos COBIT 2019 Recomendados", h2_style))
    proc_headers = ["Código", "Proceso COBIT", "Dominio", "Nivel de Urgencia", "Objetivo"]
    proc_rows = [[Paragraph(f"<b>{h}</b>", table_header) for h in proc_headers]]

    procs_data = [
        ("APO13", "Gestionar la Seguridad", "APO", "Crítica", "SGSI y protección perimetral"),
        ("DSS05", "Gestionar los Servicios de Seguridad", "DSS", "Crítica", "Despliegue de MFA y EDR"),
        ("DSS04", "Gestionar la Continuidad", "DSS", "Alta", "Plan DRP y simulacros de respaldo"),
        ("DSS02", "Gestionar Solicitudes e Incidentes", "DSS", "Alta", "Optimización de la mesa de ayuda"),
        ("MEA03", "Gestionar la Conformidad Externa", "MEA", "Media", "Cumplimiento normativo Supersalud")
    ]
    for p_id, p_name, p_dom, p_urg, p_obj in procs_data:
        proc_rows.append([
            Paragraph(p_id, table_cell),
            Paragraph(p_name, table_cell),
            Paragraph(p_dom, table_cell),
            Paragraph(p_urg, table_cell),
            Paragraph(p_obj, table_cell)
        ])

    t_procs = Table(proc_rows, colWidths=[45, 175, 45, 80, 175])
    t_procs.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E3A8A')),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F8FAFC')]),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
    ]))
    elements.append(t_procs)

    doc.build(elements)
    print(f"PDF generado exitosamente en: {pdf_path}")

def generar_docx_logistica():
    docx_path = DATA_DIR / "Plan_Estrategico_Gobernanza_TI_Logistica_Andina.docx"
    doc = docx.Document()

    # Título principal
    doc.add_heading("LOGÍSTICA Y CARGA ANDINA S.A.S.", 0)
    p_sub = doc.add_paragraph("DIAGNÓSTICO PRELIMINAR DE GOBERNANZA DE TI, CONTROL DE CAMBIOS Y GESTIÓN DE PROVEEDORES")
    p_sub.italic = True

    # Ficha del Cliente
    doc.add_heading("1. Información Corporativa y Contexto", level=1)
    doc.add_paragraph(
        "Logística y Carga Andina S.A.S. es una compañía de transporte de carga pesada, almacenamiento y distribución con operaciones "
        "en Colombia, Ecuador y Perú. Cuenta con una flota monitoreada de más de 400 vehículos de carga y centros de distribución (CEDI) "
        "en 5 ciudades principales. Actualmente la empresa experimenta un crecimiento acelerado del 30% anual en envíos de comercio electrónico, "
        "lo cual ha puesto en tensión su infraestructura tecnológica y su modelo de gobierno de TI."
    )

    # Situación actual y problemas
    doc.add_heading("2. Diagnóstico de Problemas Operacionales y de Gestión", level=1)
    doc.add_paragraph(
        "La revisión efectuada con la Gerencia General y los coordinadores de logística identificó las siguientes deficiencias principales:"
    )

    p1 = doc.add_paragraph()
    p1.add_run("• Carencia de Gobernanza y Desalineación Estratégica: ").bold = True
    p1.add_run(
        "La empresa carece de un Comité Directivo de Gobernanza de TI formal. Las inversiones tecnológicas se aprueban por urgencia "
        "comercial sin análisis de retorno de inversión (ROI), generando compras redundantes de software y costos operativos imprevistos."
    )

    p2 = doc.add_paragraph()
    p2.add_run("• Fallas Severas por Cambios sin Control en Producción: ").bold = True
    p2.add_run(
        "Se detectaron constantes caídas del sistema de tracking satelital (GPS) debido a que los desarrolladores aplican modificaciones "
        "directamente en los servidores de producción sin autorización previa de un Comité de Control de Cambios (CAB), generando retrasos en despachos."
    )

    p3 = doc.add_paragraph()
    p3.add_run("• Problemas Contractuales con Proveedores Tecnológicos: ").bold = True
    p3.add_run(
        "El proveedor principal del software WMS (Warehouse Management System) no cuenta con Acuerdos de Nivel de Servicio (SLA) "
        "firmados, lo que causa demoras de hasta 48 horas en la atención de fallas críticas sin penalidades contractuales aplicables."
    )

    p4 = doc.add_paragraph()
    p4.add_run("• Carencia de Inventario de Activos y Licenciamiento: ").bold = True
    p4.add_run(
        "No existe un inventario unificado de hardware ni control de licencias de software, existiendo riesgo de multas por uso de software "
        "no autorizado en las terminales de los conductores y coordinadores de ruta."
    )

    # Necesidades
    doc.add_heading("3. Necesidades Inmediatas del Negocio", level=1)
    doc.add_paragraph(
        "Para estabilizar la plataforma operativa y acompañar la expansión del negocio, la organización:"
    )
    doc.add_paragraph("1. Requiere conformar el Comité de Gobernanza de TI con participación del Gerente General, Financiero y de Operaciones.")
    doc.add_paragraph("2. Necesita estandarizar el procedimiento de control de cambios (CAB) antes de cualquier paso a producción.")
    doc.add_paragraph("3. Debe renegociar los contratos de terceros estableciendo SLAs medibles con penalidades claras por indisponibilidad.")
    doc.add_paragraph("4. Es prioritario definir una política de compras y gestión de activos de TI institucional.")

    # Oportunidades
    doc.add_heading("4. Oportunidades de Transformación e Innovación", level=1)
    doc.add_paragraph(
        "• Optimizar la cadena de suministro integrando un portal web B2B de autogestión de despachos para los clientes corporativos."
    )
    doc.add_paragraph(
        "• Automatizar el seguimiento de la flota mediante analítica de datos e inteligencia artificial para predecir rutas óptimas y consumo de combustible."
    )
    doc.add_paragraph(
        "• Beneficio económico proyectado: Reducción del 25% en tiempos de espera en muelles y disminución de costos operativos en TI."
    )

    # Tabla de Procesos COBIT
    doc.add_heading("5. Procesos COBIT 2019 de Intervención Prioritaria", level=1)
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Light Shading Accent 1'
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = 'Código'
    hdr_cells[1].text = 'Proceso COBIT'
    hdr_cells[2].text = 'Dominio'
    hdr_cells[3].text = 'Acción Propuesta'

    row_data = [
        ("EDM01", "Asegurar el Marco de Gobernanza", "EDM", "Crear estatuto del comité de TI"),
        ("APO02", "Gestionar la Estrategia", "APO", "Alinear plan PETI con metas de transporte"),
        ("APO09", "Gestionar los SLAs", "APO", "Formalizar acuerdos de servicio y niveles de soporte"),
        ("APO10", "Gestionar Proveedores", "APO", "Auditar contratos de terceros y proveedores de GPS"),
        ("BAI06", "Gestionar los Cambios de TI", "BAI", "Digitalizar flujo de aprobación en Comité CAB"),
        ("BAI09", "Gestionar los Activos", "BAI", "Levantar inventario físico y licencias"),
        ("APO14", "Gestionar los Datos", "APO", "Gobernanza de datos de telemetría y ruteo")
    ]

    for c_id, c_name, c_dom, c_act in row_data:
        row_cells = table.add_row().cells
        row_cells[0].text = c_id
        row_cells[1].text = c_name
        row_cells[2].text = c_dom
        row_cells[3].text = c_act

    doc.save(str(docx_path))
    print(f"DOCX generado exitosamente en: {docx_path}")

def generar_archivos_texto():
    # También generamos versiones en texto plano TXT para máxima compatibilidad o copia directa
    txt1 = DATA_DIR / "Auditoria_Ciberseguridad_Hospital_San_Rafael.txt"
    txt2 = DATA_DIR / "Plan_Gobernanza_TI_Logistica_Andina.txt"

    with open(txt1, "w", encoding="utf-8") as f:
        f.write("""HOSPITAL UNIVERSITARIO SAN RAFAEL E.S.E. - INFORME DE AUDITORIA DE TI
Sector: Salud y Asistencia Medica | Fecha: Octubre 2026

RESUMEN EJECUTIVO:
El Hospital San Rafael atiende a mas de 250,000 pacientes y custodia historias clinicas electronicas con datos sensibles de salud.
La auditoria evaluo la seguridad de la informacion, procedimientos de copias de seguridad (backup), gestion de incidentes y cumplimiento regulatorio.

PROBLEMAS Y DEFICIENCIAS DETECTADAS:
- Problema critico en ciberseguridad e identidades: La entidad carece de autenticacion multifactor (MFA) en los correos y accesos remotos VPN, generando alto riesgo de ransomware y robo de datos de pacientes.
- Fallas en copias de seguridad (backup): Las copias de respaldo de la base de datos clinica no tienen pruebas de restauracion y no existe sitio alterno.
- Ausencia de plan de continuidad (DRP): No existe un plan de contingencia ante caidas de sistemas en salas de urgencias y cirugia.
- Tiempos de retraso en la mesa de ayuda: Retrasos del 45% en la resolucion de incidentes de red que afectan la dispensacion de medicamentos.

NECESIDADES PRIORITARIAS:
1. Se requiere formalizar e implementar la politica de seguridad de la informacion (SGSI) alineada a COBIT APO13 y DSS05.
2. La entidad necesita establecer acuerdos de nivel de servicio (SLA) con el proveedor de software clinico para tiempos de respuesta.
3. Se debe ejecutar simulacros de recuperacion de backups para garantizar la continuidad operativa.
4. Es fundamental capacitar al personal medico en prevencion de ciberamenazas y phishing.

OPORTUNIDADES DE MEJORA:
- Oportunidad de modernizar copias de seguridad en la nube con almacenamiento inmutable contra malware.
- Automatizar el monitoreo de seguridad perimetral mediante alertas tempranas.
- Potencial de mejora en el Service Desk con herramientas de autoservicio de contrasenas para medicos de turno.
""")

    with open(txt2, "w", encoding="utf-8") as f:
        f.write("""LOGISTICA Y CARGA ANDINA S.A.S. - DIAGNOSTICO DE GOBERNANZA DE TI
Sector: Transporte y Cadena de Suministro | Fecha: Octubre 2026

RESUMEN EJECUTIVO:
Compania lider en distribucion de carga con mas de 400 vehiculos monitoreados y 5 centros de distribucion logisticos.
Presenta una tasa de crecimiento del 30% anual en comercio electronico que requiere fortalecer su gestion de TI.

PROBLEMAS IDENTIFICADOS:
- Desalineacion y falta de gobernanza: La empresa carece de un Comite Directivo de Gobernanza de TI formal. Las inversiones se realizan de manera reactiva sin calculo de retorno financiero (ROI).
- Fallas constantes por cambios no controlados: Caidas del software de monitoreo GPS debido a cambios aplicados en produccion sin aprobacion del comite CAB.
- Problemas con proveedores tecnologicos: El proveedor del sistema de almacenes (WMS) no tiene acuerdos de nivel de servicio (SLAs) formalizados, provocando retrasos en despachos.
- Inadecuado control de inventario de activos y software sin licencias centralizadas.

NECESIDADES PRIORITARIAS:
1. La organizacion requiere conformar el Comite de Gobernanza de TI con participacion de la gerencia general.
2. Se necesita estandarizar el proceso de gestion de cambios (CAB) antes de cualquier paso a produccion.
3. Se debe renegociar los contratos de proveedores de tecnologia incluyendo SLAs obligatorios.
4. Es prioritario levantar el inventario de hardware y licencias para evitar sanciones legales.

OPORTUNIDADES DE OPTIMIZACION:
- Optimizar la cadena de suministro habilitando un portal web de autogestion para clientes corporativos.
- Automatizar el ruteo de camiones mediante analitica avanzada de datos de telemetria para ahorro de combustible.
- Beneficio estimado: Reduccion del 25% en tiempos de espera en muelles logisticos.
""")

    print(f"Archivos de texto generados en: {txt1} y {txt2}")

if __name__ == "__main__":
    generar_pdf_salud()
    generar_docx_logistica()
    generar_archivos_texto()
    print("¡Todos los documentos de prueba fueron generados con éxito!")
